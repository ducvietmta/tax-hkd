import json, html, zipfile, shutil
from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.pagesizes import A4
from pypdf import PdfReader
import pypdfium2 as pdfium
from PIL import Image, ImageDraw

ROOT=Path(__file__).resolve().parents[1]
data=json.loads((ROOT/'shared/websitePolicies.json').read_text(encoding='utf-8'))
out=ROOT/'outputs/website-policies-final-2026-10-23';public=ROOT/'public/policies'
out.mkdir(parents=True,exist_ok=True);public.mkdir(parents=True,exist_ok=True)
owner=data['operator'];assert all(owner.get(k) for k in ['name','taxId','address','email','phone'])
identity=[f"Chủ sở hữu {owner['name']} · Mã số thuế {owner['taxId']}",owner['address'],f"Email {owner['email']} · Hotline/Zalo {owner['phone']}",f"Website {data['website']} · Cập nhật 23/10/2026"]
pdfmetrics.registerFont(TTFont('Arial','C:/Windows/Fonts/arial.ttf'));pdfmetrics.registerFont(TTFont('ArialB','C:/Windows/Fonts/arialbd.ttf'))
styles={
 'title':ParagraphStyle('title',fontName='ArialB',fontSize=16,leading=22,spaceAfter=12),
 'meta':ParagraphStyle('meta',fontName='Arial',fontSize=9,leading=13,spaceAfter=4),
 'head':ParagraphStyle('head',fontName='ArialB',fontSize=11,leading=16,spaceBefore=11,spaceAfter=5,keepWithNext=True),
 'body':ParagraphStyle('body',fontName='Arial',fontSize=10.5,leading=16,spaceAfter=8,allowWidows=0,allowOrphans=0),
}
def footer(canvas,doc):
 canvas.setFont('Arial',8);canvas.drawString(48,25,'TAX-HKD · Chính sách dịch vụ');canvas.drawRightString(A4[0]-48,25,f'Trang {doc.page}')
paths=[];thumbs=[];report=[]
for policy in data['policies']:
 stem=f"{policy['number']:02}-{policy['id']}";title=policy['title'];doc=Document();sec=doc.sections[0]
 sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=sec.bottom_margin=Cm(1.8);sec.left_margin=sec.right_margin=Cm(1.9)
 for name in ['Normal','Title','Heading 1','Heading 2']:
  style=doc.styles[name];style.font.name='Arial';style.font.color.rgb=RGBColor(0,0,0)
 doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.line_spacing=1.12
 doc.styles['Normal'].paragraph_format.space_after=Pt(7);doc.styles['Title'].font.size=Pt(18)
 doc.styles['Heading 1'].font.size=Pt(12);doc.styles['Heading 1'].paragraph_format.keep_with_next=True
 doc.add_paragraph(title,style='Title')
 story=[Paragraph(escape(title),styles['title'])]
 for line in identity:
  para=doc.add_paragraph(line);para.paragraph_format.space_after=Pt(3)
  for run in para.runs:run.font.size=Pt(9)
  story.append(Paragraph(escape(line),styles['meta']))
 story.append(Spacer(1,10))
 for i,section in enumerate(policy['sections'],1):
  heading=f"{i}. {section['heading']}";doc.add_heading(heading,level=1);story.append(Paragraph(escape(heading),styles['head']))
  for text in section['paragraphs']:doc.add_paragraph(text);story.append(Paragraph(escape(text),styles['body']))
 doc.sections[0].footer.paragraphs[0].text='TAX-HKD · '+title
 docx_path=out/(stem+'.docx');pdf_path=out/(stem+'.pdf');doc.save(docx_path)
 SimpleDocTemplate(str(pdf_path),pagesize=A4,rightMargin=48,leftMargin=48,topMargin=42,bottomMargin=43).build(story,onFirstPage=footer,onLaterPages=footer)
 text='\n'.join(page.extract_text() for page in PdfReader(pdf_path).pages)
 for field in [owner['name'],owner['taxId'],owner['email']]:assert field in text
 assert ' '.join(title.split()) in ' '.join(text.split());assert len(Document(docx_path).paragraphs)>len(policy['sections'])
 rendered=pdfium.PdfDocument(str(pdf_path))
 for j,page in enumerate(rendered):
  image=page.render(scale=.7).to_pil().convert('RGB');image.thumbnail((300,430));thumbs.append((stem+f' p{j+1}',image.copy()))
 report.append({'file':stem,'pages':len(rendered),'characters':len(text)})
 files=[docx_path,pdf_path]
 contents=''.join('<section><h2>'+html.escape(str(i)+'. '+s['heading'])+'</h2>'+''.join('<p>'+html.escape(p)+'</p>' for p in s['paragraphs'])+'</section>' for i,s in enumerate(policy['sections'],1))
 page='<!doctype html><html lang="vi"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+html.escape(title)+' | TAX-HKD</title><style>body{font:16px/1.75 Arial;margin:40px auto;padding:0 24px;max-width:850px;color:#17212b}h1{font-size:28px;line-height:1.35}h2{font-size:19px;margin-top:28px}small{color:#536170}a{color:#08775c}nav{display:flex;gap:16px;flex-wrap:wrap}@media print{nav{display:none}}</style><nav><a href="/">TAX-HKD</a><a href="'+stem+'.pdf">Tải PDF</a><a href="'+stem+'.docx">Tải Word</a><a href="TAX-HKD-6-chinh-sach.zip">Tải cả bộ</a></nav><h1>'+html.escape(title)+'</h1><small>'+'<br>'.join(map(html.escape,identity))+'</small>'+contents+'</html>'
 html_path=out/(stem+'.html');html_path.write_text(page,encoding='utf-8');files.append(html_path)
 for f in files:shutil.copy2(f,public/f.name);paths.append(f)
with zipfile.ZipFile(out/'TAX-HKD-6-chinh-sach.zip','w',zipfile.ZIP_DEFLATED) as z:
 for file in paths:z.write(file,file.name)
shutil.copy2(out/'TAX-HKD-6-chinh-sach.zip',public/'TAX-HKD-6-chinh-sach.zip')
sheet=Image.new('RGB',(960,480*((len(thumbs)+2)//3)),'#dddddd');draw=ImageDraw.Draw(sheet)
for i,(label,pic) in enumerate(thumbs):x=(i%3)*320+10;y=(i//3)*480;draw.text((x,y+5),label,fill='black');sheet.paste(pic,(x,y+25))
sheet.save(out/'visual-review.png');(out/'verification.json').write_text(json.dumps(report,indent=2),encoding='utf-8');print(json.dumps(report))
